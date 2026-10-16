from pathlib import Path
import re, sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(r'C:\Users\perez\OneDrive\Escritorio\plataforma-c2c\plataforma-c2c')
source = root / 'docs' / 'TALLER-1.md'
out = root / 'docs' / 'TALLER-1.docx'
text = source.read_text(encoding='utf-8').replace('\r\n','\n')
lines = text.splitlines()
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.62); sec.bottom_margin = Inches(.62)
sec.left_margin = Inches(.72); sec.right_margin = Inches(.72)
sec.header_distance = Inches(.3); sec.footer_distance = Inches(.3)

sage = '657A4B'; ink = '252B23'; muted='657064'; pale='EAF0E1'; codebg='F2F4EF'
def set_font(style, name, size, color=ink, bold=False):
    style.font.name=name; style.font.size=Pt(size); style.font.color.rgb=RGBColor.from_string(color); style.font.bold=bold
    style._element.rPr.rFonts.set(qn('w:ascii'),name); style._element.rPr.rFonts.set(qn('w:hAnsi'),name)
set_font(doc.styles['Normal'],'Aptos',9.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
set_font(doc.styles['Title'],'Aptos Display',25,ink,True)
set_font(doc.styles['Subtitle'],'Aptos',12,sage)
set_font(doc.styles['Heading 1'],'Aptos Display',17,ink,True)
set_font(doc.styles['Heading 2'],'Aptos Display',13,sage,True)
set_font(doc.styles['Heading 3'],'Aptos',10,ink,True)

header=sec.header.paragraphs[0]
header.text='C2C MARKET  |  TALLER 1 · SISTEMA COMPLEJO'
header.style='Caption'; header.runs[0].font.color.rgb=RGBColor.from_string(muted)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('C2C Market  ·  '); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(muted)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)
def add_link(paragraph, label, url):
    part=paragraph.part; rid=part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rid); run=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),sage); pr.append(color)
    u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); pr.append(u); run.append(pr)
    t=OxmlElement('w:t'); t.text=label; run.append(t); h.append(run); paragraph._p.append(h)
def inline(paragraph, value):
    pattern=re.compile(r'(\[[^\]]+\]\([^)]+\)|`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)')
    pos=0
    for m in pattern.finditer(value):
        if m.start()>pos: paragraph.add_run(value[pos:m.start()])
        token=m.group(0)
        if token.startswith('['):
            z=re.match(r'\[([^\]]+)\]\(([^)]+)\)',token); add_link(paragraph,z.group(1),z.group(2))
        elif token.startswith('`'):
            rr=paragraph.add_run(token[1:-1]); rr.font.name='Consolas'; rr.font.size=Pt(8); rr.font.color.rgb=RGBColor.from_string('45513F')
        elif token.startswith('**'):
            rr=paragraph.add_run(token[2:-2]); rr.bold=True
        else:
            rr=paragraph.add_run(token[1:-1]); rr.italic=True
        pos=m.end()
    if pos<len(value): paragraph.add_run(value[pos:])

# Formal, editable academic cover.
title=doc.add_paragraph(style='Title'); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
title.add_run('TALLER 1 — SISTEMA COMPLEJO\nCONSTRUCCIÓN Y DESARROLLO DE SOFTWARE')
sub=doc.add_paragraph(style='Subtitle'); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
sub.add_run('Arquitectura C4, endpoints REST, frontend web y despliegue en Vercel de C2C Market')
sp=doc.add_paragraph(); sp.paragraph_format.space_before=Pt(24); sp.alignment=WD_ALIGN_PARAGRAPH.CENTER
rr=sp.add_run('C2C MARKET'); rr.bold=True; rr.font.size=Pt(16); rr.font.color.rgb=RGBColor.from_string(sage)
sp=doc.add_paragraph(); sp.alignment=WD_ALIGN_PARAGRAPH.CENTER; sp.paragraph_format.space_before=Pt(16)
sp.add_run('Lo que buscas, alguien lo tiene').italic=True
for label in ['Estudiante','Docente','Asignatura','Institución','Fecha']:
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(1.2); p.paragraph_format.space_before=Pt(8)
    a=p.add_run(label+':  '); a.bold=True; a.font.color.rgb=RGBColor.from_string(sage); p.add_run('________________________________________')
doc.add_page_break()

# Lightweight Markdown-to-Word conversion with styled code and tables.
i=0; in_code=False; code=[]; lang=''
while i<len(lines):
    line=lines[i]
    if line.startswith('```'):
        if not in_code:
            in_code=True; lang=line[3:].strip(); code=[]
        else:
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.1); p.paragraph_format.right_indent=Inches(.08)
            p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
            pPr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),codebg); pPr.append(shd)
            borders=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'nil'); borders.append(bottom); pPr.append(borders)
            for idx,codeline in enumerate(code):
                if idx: p.add_run().add_break()
                rr=p.add_run(codeline if codeline else ' '); rr.font.name='Consolas'; rr.font.size=Pt(6.7); rr.font.color.rgb=RGBColor.from_string('30382F')
            in_code=False; lang=''; code=[]
        i+=1; continue
    if in_code:
        code.append(line); i+=1; continue
    if line.startswith('|') and i+1<len(lines) and re.match(r'^\|?\s*:?-{3,}',lines[i+1]):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            vals=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-{3,}:?',x or '-') for x in vals): rows.append(vals)
            i+=1
        if rows:
            cols=max(map(len,rows)); table=doc.add_table(rows=1,cols=cols); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
            for ri,row in enumerate(rows):
                cells=table.rows[0].cells if ri==0 else table.add_row().cells
                for ci in range(cols):
                    cell=cells[ci]; cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text=''; pp=cell.paragraphs[0]; pp.paragraph_format.space_after=Pt(1)
                    inline(pp,row[ci] if ci<len(row) else '')
                    for run in pp.runs: run.font.size=Pt(7.4)
                    if ri==0:
                        shade(cell,pale)
                        for run in pp.runs: run.bold=True; run.font.color.rgb=RGBColor.from_string(ink)
            set_repeat_table_header(table.rows[0]); doc.add_paragraph().paragraph_format.space_after=Pt(1)
        continue
    if not line.strip(): i+=1; continue
    if line.startswith('# '):
        p=doc.add_paragraph(style='Heading 1'); inline(p,line[2:])
    elif line.startswith('## '):
        p=doc.add_paragraph(style='Heading 1'); p.paragraph_format.keep_with_next=True
        if 'Anexo' in line: p.paragraph_format.page_break_before=True
        inline(p,line[3:])
    elif line.startswith('### '):
        p=doc.add_paragraph(style='Heading 2'); inline(p,line[4:])
    elif line.startswith('- '):
        p=doc.add_paragraph(style='List Bullet'); inline(p,line[2:])
    elif re.match(r'^\d+\. ',line):
        p=doc.add_paragraph(style='List Number'); inline(p,re.sub(r'^\d+\. ','',line))
    else:
        p=doc.add_paragraph(); inline(p,line)
    i+=1

doc.core_properties.title='Taller 1 Sistema Complejo C2C Market'
doc.core_properties.subject='Arquitectura C4, API REST, frontend y despliegue'
doc.core_properties.author=''
out.parent.mkdir(parents=True,exist_ok=True); doc.save(out)
print(out)
